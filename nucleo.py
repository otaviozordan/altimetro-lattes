"""Nucleo do Altimetro: grupos, download dos curriculos e scriptLattes, leitura
dos dados, metricas, tabelas, graficos, texto e Excel. A janela fica em
altimetro.py; o download e o scriptLattes rodam num processo separado
(altimetro.py --baixar / --gerar), para a janela nao travar.

Cada grupo fica em grupos/<nome>/: membros.json, cache/ (curriculos),
saida/ (o que o scriptLattes gera) e relatorios/ (relatorios salvos).
"""
import collections
import datetime
import glob
import io
import json
import os
import platform
import queue
import re
import shutil
import subprocess
import sys
import tempfile
import threading
import unicodedata
import urllib.request
import webbrowser
import math

import marca

AQUI = os.path.dirname(os.path.abspath(__file__))
SCRIPTLATTES = os.path.join(AQUI, "scriptLattes")
MODELO_CONFIG = os.path.join(AQUI, "scriptLattes.config.modelo")
GRUPOS = os.path.join(AQUI, "grupos")
PREFERENCIAS = os.path.join(AQUI, "preferencias.json")

CATEGORIAS = ["Docentes", "Pesquisadores", "Pós-doutorandos", "Doutorandos", "Mestrandos",
              "Graduandos", "Egressos", "Colaboradores externos"]


# ------------------------------------------------------------------ utilidades

def norm(s):
    s = unicodedata.normalize("NFD", str(s or ""))
    return "".join(c for c in s if not unicodedata.combining(c)).lower().strip()


def chave_texto(s, n=150):
    return re.sub(r"[^a-z0-9]+", "", norm(s))[:n]


def ano_int(x):
    m = re.search(r"(19|20)\d{2}", str(x if x is not None else ""))
    return int(m.group()) if m else None


def doi_de(x):
    m = re.search(r"10\.\d{4,9}/\S+", str(x or ""), re.I)
    return m.group().lower().rstrip(".,;") if m else ""


def num(n):
    return f"{n:,}".replace(",", ".")


def qt(n, sing, plur):
    return f"{num(n)} {sing if n == 1 else plur}"


def lista_pt(partes):
    partes = [p for p in partes if p]
    if len(partes) <= 1:
        return "".join(partes)
    return ", ".join(partes[:-1]) + " e " + partes[-1]


def ler_texto(caminho):
    dados = open(caminho, "rb").read()
    for cod in ("utf-8", "cp1252"):
        try:
            return dados.decode(cod)
        except UnicodeDecodeError:
            pass
    return dados.decode("latin-1")


def ler_json(caminho, padrao):
    try:
        with open(caminho, encoding="utf-8") as f:
            return json.load(f)
    except (OSError, ValueError):
        return padrao


def gravar_json(caminho, dados):
    tmp = caminho + ".tmp"
    with open(tmp, "w", encoding="utf-8") as f:
        json.dump(dados, f, ensure_ascii=False, indent=1)
    os.replace(tmp, caminho)


def abrir(caminho):
    """Abre um arquivo ou pasta no programa padrao (Windows, macOS ou Linux)."""
    if platform.system() == "Windows":
        os.startfile(caminho)  # noqa: S606
    elif platform.system() == "Darwin":
        subprocess.Popen(["open", caminho])
    else:
        subprocess.Popen(["xdg-open", caminho])


FOTO_URL = "http://servicosweb.cnpq.br/wspessoa/servletrecuperafoto?tipo=1&id="


def id_foto(html):
    m = re.search(r"servletrecuperafoto\?tipo=1&(?:amp;)?id=(K\w+)", html)
    return m.group(1) if m else ""


def salvar_foto(dados, destino):
    """Grava a foto em JPEG quadrado (centro), 320 px. False se for a imagem padrao do Lattes."""
    from PIL import Image
    im = Image.open(io.BytesIO(dados)).convert("RGB")
    if min(im.size) < 120:
        return False
    lado = min(im.size)
    x, y = (im.width - lado) // 2, max(0, (im.height - lado) // 4)  # rosto costuma ficar no alto
    im = im.crop((x, y, x + lado, y + lado)).resize((320, 320), Image.LANCZOS)
    os.makedirs(os.path.dirname(destino), exist_ok=True)
    im.save(destino, "JPEG", quality=88)
    return True


def baixar_foto(g, lid, html=None):
    """Foto do Lattes do membro para fotos/<id>.jpg. Sem foto no Lattes: grava
    fotos/<id>.sem, para nao tentar de novo a cada atualizacao."""
    if html is None:
        html = ler_texto(g.cv(lid)) if os.path.exists(g.cv(lid)) else ""
    kid = id_foto(html)
    sem = os.path.join(g.fotos, lid + ".sem")
    if not kid:
        os.makedirs(g.fotos, exist_ok=True)
        open(sem, "w").close()
        return False
    req = urllib.request.Request(FOTO_URL + kid, headers={"User-Agent": "Mozilla/5.0"})
    with urllib.request.urlopen(req, timeout=60) as r:  # noqa: S310 - endereco fixo do CNPq
        dados = r.read()
    try:
        ok = salvar_foto(dados, g.foto_arquivo(lid))
    except Exception:  # noqa: BLE001 - imagem invalida
        ok = False
    if not ok:
        os.makedirs(g.fotos, exist_ok=True)
        open(sem, "w").close()
    elif os.path.exists(sem):
        os.remove(sem)
    return ok


def nome_no_cv(html):
    m = re.search(r'<h2 class="nome"[^>]*>\s*([^<]+)', html)
    return re.sub(r"\s+", " ", m.group(1)).strip() if m else ""


# ---------------------------------------------------------------------- grupo

class Grupo:
    """Uma pasta em grupos/: membros.json, cache/ (CVs) e saida/ (scriptLattes)."""

    def __init__(self, pasta):
        self.pasta = pasta
        self.arquivo = os.path.join(pasta, "membros.json")
        self.cache = os.path.join(pasta, "cache")
        self.saida = os.path.join(pasta, "saida")
        d = ler_json(self.arquivo, {})
        self.nome = d.get("nome") or os.path.basename(pasta)
        self.membros = d.get("membros", [])  # [{"id","nome","categoria","incluir"}]

    @staticmethod
    def criar(nome):
        base = re.sub(r"[^\w\- ]+", "", nome).strip() or "grupo"
        pasta, n = os.path.join(GRUPOS, base), 2
        while os.path.exists(pasta):
            pasta, n = os.path.join(GRUPOS, f"{base} {n}"), n + 1
        os.makedirs(pasta)
        g = Grupo(pasta)
        g.nome = nome
        g.salvar()
        return g

    @staticmethod
    def todos():
        os.makedirs(GRUPOS, exist_ok=True)
        return [Grupo(p) for p in sorted(glob.glob(os.path.join(GRUPOS, "*"))) if os.path.isdir(p)]

    def salvar(self):
        os.makedirs(self.pasta, exist_ok=True)
        gravar_json(self.arquivo, {"nome": self.nome, "membros": self.membros})

    def membro(self, lid):
        for m in self.membros:
            if m["id"] == lid:
                return m
        return None

    def cv(self, lid):
        return os.path.join(self.cache, lid)

    def cv_data(self, lid):
        c = self.cv(lid)
        return datetime.date.fromtimestamp(os.path.getmtime(c)) if os.path.exists(c) else None

    @property
    def fotos(self):
        return os.path.join(self.pasta, "fotos")

    def foto_arquivo(self, lid):
        return os.path.join(self.fotos, lid + ".jpg")

    def foto(self, lid):
        c = self.foto_arquivo(lid)
        return c if os.path.exists(c) else None

    @property
    def relatorios(self):
        return os.path.join(self.pasta, "relatorios")

    def tem_relatorio(self):
        return os.path.exists(os.path.join(self.saida, "index.html"))


# ------------------------------------- download e scriptLattes (outro processo)

def avisar(*partes):
    print("@@", *partes, flush=True)


def cv_valido(html, lid):
    return "Curr" in html and lid in html and "captcha" not in html.lower()[:2000]


def cmd_baixar(pasta, so_faltam):
    """Baixa os CVs com o baixaLattes do scriptLattes (Chrome headless)."""
    sys.path.insert(0, SCRIPTLATTES)
    from scriptLattes.baixaLattes import baixaCVLattes
    g = Grupo(pasta)
    os.makedirs(g.cache, exist_ok=True)
    lista = [m for m in g.membros if not (so_faltam and os.path.exists(g.cv(m["id"])))]
    if not lista:
        avisar("fim", "ok", "Todos os currículos já estavam baixados.")
        return
    # Uma subpasta nova por CV: o baixaCVLattes pula o que ja existe no destino,
    # e no OneDrive apagar e recriar a mesma pasta costuma falhar.
    raiz = os.path.join(pasta, "_baixando")
    os.makedirs(raiz, exist_ok=True)
    ok, falhas = 0, []
    for k, m in enumerate(lista, 1):
        nome = m.get("nome") or m["id"]
        avisar("prog", k - 1, len(lista), "Baixando " + nome)
        print(f"[{k}/{len(lista)}] {nome}", flush=True)
        tmp = tempfile.mkdtemp(dir=raiz)
        try:
            baixaCVLattes(m["id"], tmp)
            html = ler_texto(os.path.join(tmp, m["id"]))
            if not cv_valido(html, m["id"]):
                raise ValueError("a página baixada não é o currículo")
        except Exception as e:  # noqa: BLE001 - um CV com problema nao para os outros
            falhas.append(nome)
            print("  falhou:", (str(e).splitlines() or [type(e).__name__])[0], flush=True)
            continue
        with open(g.cv(m["id"]), "w", encoding="utf-8") as f:
            f.write(html)
        ok += 1
        if not os.path.exists(os.path.join(g.fotos, m["id"] + ".usuario")):  # foto escolhida a mao fica
            try:
                baixar_foto(g, m["id"], html)
            except Exception as e:  # noqa: BLE001 - sem foto, segue
                print("  foto nao baixou:", type(e).__name__, flush=True)
        if not m.get("nome") or m["nome"] == m["id"]:
            g2 = Grupo(pasta)  # relido: a janela pode ter mudado algo enquanto isso
            mm = g2.membro(m["id"])
            if mm is not None:
                mm["nome"] = nome_no_cv(html) or mm.get("nome", "")
                g2.salvar()
    shutil.rmtree(raiz, ignore_errors=True)
    msg = f"{ok} de {len(lista)} currículos baixados."
    if falhas:
        msg += " Falharam: " + ", ".join(falhas) + "."
    avisar("prog", len(lista), len(lista), msg)
    avisar("fim", "ok" if ok or not falhas else "erro", msg)


def cmd_gerar(pasta):
    """Roda o scriptLattes com os CVs que estao no cache."""
    g = Grupo(pasta)
    presentes = [m for m in g.membros if os.path.exists(g.cv(m["id"]))]
    if not presentes:
        avisar("fim", "erro", "Nenhum currículo baixado ainda.")
        return
    faltam_fotos = [m for m in presentes if not g.foto(m["id"]) and not os.path.exists(os.path.join(g.fotos, m["id"] + ".sem"))]
    for k, m in enumerate(faltam_fotos, 1):
        avisar("prog", 0, 1, f"Buscando fotos do Lattes ({k}/{len(faltam_fotos)})…")
        try:
            baixar_foto(g, m["id"])
        except Exception as e:  # noqa: BLE001
            print("  foto nao baixou:", m.get("nome") or m["id"], type(e).__name__, flush=True)
    avisar("prog", 0, 1, f"Rodando o scriptLattes com {len(presentes)} currículos…")
    lista = os.path.join(pasta, "grupo.list")
    with open(lista, "w", encoding="utf-8") as f:
        for m in presentes:
            nome = (m.get("nome") or m["id"]).replace(",", " ")
            f.write(f"{m['id']} , {nome} , , {m.get('categoria') or 'Membros'}\n")
    termos = os.path.join(pasta, "termos.txt")
    open(termos, "a").close()
    if os.path.isdir(g.saida):
        shutil.rmtree(g.saida, ignore_errors=True)
    config = os.path.join(pasta, "grupo.config")
    with open(config, "w", encoding="utf-8") as f:
        f.write(open(MODELO_CONFIG, encoding="utf-8").read()
                .replace("{NOME}", g.nome).replace("{LISTA}", lista)
                .replace("{SAIDA}", g.saida + os.sep).replace("{CACHE}", g.cache + os.sep)
                .replace("{TERMOS}", termos).replace("{ANO}", str(datetime.date.today().year)))
    env = dict(os.environ, PYTHONUTF8="1", PYTHONIOENCODING="utf-8")
    r = subprocess.run([sys.executable, os.path.join(SCRIPTLATTES, "scriptLattes.py"), config],
                       cwd=SCRIPTLATTES, env=env, stdout=sys.stdout, stderr=subprocess.STDOUT)
    if r.returncode != 0 or not g.tem_relatorio():
        avisar("fim", "erro", "O scriptLattes falhou; veja o registro.")
        return
    # O nome completo vem do proprio CV: completa quem foi cadastrado so pelo ID.
    g2 = Grupo(pasta)
    for arq in glob.glob(os.path.join(g.saida, "json", "*.json")):
        info = ler_json(arq, {}).get("informacoes_pessoais", {})
        mm = g2.membro(str(info.get("id_lattes", "")))
        if mm is not None and info.get("nome_completo") and (not mm.get("nome") or mm["nome"] == mm["id"]):
            mm["nome"] = info["nome_completo"]
    g2.salvar()
    avisar("prog", 1, 1, "Relatório gerado.")
    avisar("fim", "ok", f"Estatísticas geradas com {len(presentes)} currículos.")


# -------------------------------------------------------- leitura dos dados

GRUPOS_CAT = ["Produção bibliográfica", "Produção técnica", "Orientações concluídas",
              "Orientações em andamento", "Projetos", "Outras atividades"]

BIB = [
    ("artigos", "Artigos em periódicos", "artigos_periodicos", ("revista",)),
    ("aceitos", "Artigos aceitos para publicação", "artigos_aceitos", ("revista",)),
    ("livros", "Livros publicados", "livros_publicados", ("editora",)),
    ("capitulos", "Capítulos de livros", "capitulos_livros", ("titulo_livro", "editora")),
    ("completos", "Trabalhos completos em eventos", "trabalhos_completos_congressos", ("evento",)),
    ("expandidos", "Resumos expandidos em eventos", "resumos_expandidos", ("evento",)),
    ("resumos", "Resumos em eventos", "resumos_congressos", ("evento",)),
    ("apresentacoes", "Apresentações de trabalho", "apresentacoes_trabalhos", ("evento",)),
    ("jornais", "Textos em jornais e revistas", "textos_jornais", ("jornal", "revista", "veiculo")),
    ("outras_bib", "Outras produções bibliográficas", "outras_producoes", ("natureza", "editora")),
]
NIVEIS = [
    ("posdoc", "Pós-doutorado", "pos_doutorado"),
    ("doutorado", "Doutorado", "doutorado"),
    ("mestrado", "Mestrado", "mestrado"),
    ("especializacao", "Especialização", "especializacao"),
    ("tcc", "Trabalho de conclusão de curso", "tcc"),
    ("ic", "Iniciação científica", "iniciacao_cientifica"),
    ("outras", "Outras orientações", "outros"),
]

CATS = collections.OrderedDict()
for _k, _r, _j, _v in BIB:
    CATS[_k] = (_r, "Produção bibliográfica")
CATS["software"] = ("Softwares", "Produção técnica")
CATS["patentes"] = ("Patentes e registros", "Produção técnica")
CATS["tecnica"] = ("Outras produções técnicas", "Produção técnica")
CATS["artistica"] = ("Produção artística", "Produção técnica")
for _k, _r, _j in NIVEIS:
    CATS["oc_" + _k] = (_r, "Orientações concluídas")
for _k, _r, _j in NIVEIS:
    CATS["oa_" + _k] = (_r, "Orientações em andamento")
CATS["pj_pesquisa"] = ("Projetos de pesquisa", "Projetos")
CATS["pj_extensao"] = ("Projetos de extensão", "Projetos")
CATS["pj_desenvolvimento"] = ("Projetos de desenvolvimento", "Projetos")
CATS["pj_outros"] = ("Outros projetos", "Projetos")
CATS["premios"] = ("Prêmios e títulos", "Outras atividades")
CATS["bancas"] = ("Participações em bancas", "Outras atividades")
CATS["ev_part"] = ("Participações em eventos", "Outras atividades")
CATS["ev_org"] = ("Organização de eventos", "Outras atividades")

PRODUCAO = [k for k, (_r, gr) in CATS.items() if gr in ("Produção bibliográfica", "Produção técnica")]
OC = ["oc_" + k for k, _r, _j in NIVEIS]
OA = ["oa_" + k for k, _r, _j in NIVEIS]
PJ = ["pj_pesquisa", "pj_extensao", "pj_desenvolvimento", "pj_outros"]


class Item:
    __slots__ = ("cat", "ano", "fim", "titulo", "pessoas", "veiculo", "doi", "extra", "coord", "membros")

    def __init__(self, cat, ano, titulo, pessoas="", veiculo="", doi="", extra="", fim=None, coord=()):
        self.cat, self.ano, self.fim = cat, ano, fim
        self.titulo = re.sub(r"\s+", " ", str(titulo or "")).strip()
        self.pessoas = re.sub(r"\s+", " ", str(pessoas or "")).strip()
        self.veiculo = re.sub(r"\s+", " ", str(veiculo or "")).strip()
        self.doi, self.extra = doi, extra
        self.coord = set(coord)
        self.membros = set()

    def chave(self, dono):
        if self.cat in ("premios",):
            return (self.cat, dono, chave_texto(self.titulo))
        if self.cat.startswith(("oc_", "oa_")):
            return (self.cat, chave_texto(self.titulo, 100), chave_texto(self.pessoas, 40))
        if self.cat in PJ:
            return ("pj", chave_texto(self.titulo))
        if self.cat == "bancas":
            return (self.cat, chave_texto(self.pessoas, 40), self.ano, self.extra, chave_texto(self.titulo, 60) if not self.pessoas else "")
        if self.doi:
            return ("doi", self.doi)
        return (self.cat, chave_texto(self.titulo), self.ano)


def _primeiro(d, campos):
    for c in campos:
        if d.get(c):
            return d[c]
    return ""


def itens_do_json(d):
    out = []
    pb = d.get("producao_bibliografica") or {}
    for cat, _r, chave, campos in BIB:
        for x in pb.get(chave) or []:
            out.append(Item(cat, ano_int(x.get("ano")), x.get("titulo"), x.get("autores"), _primeiro(x, campos), doi_de(x.get("doi"))))
    pt = d.get("producao_tecnica") or {}
    for chave, lista in pt.items():
        cat = "software" if chave.startswith("softwares") else "tecnica"
        for x in lista or []:
            out.append(Item(cat, ano_int(x.get("ano")), x.get("titulo"), x.get("autores"), x.get("tipo", ""), extra=x.get("tipo", "")))
    for chave, lista in (d.get("patentes_registros") or {}).items():
        for x in lista or []:
            out.append(Item("patentes", ano_int(x.get("ano")), x.get("titulo"), x.get("autores"), extra=chave.replace("_", " ")))
    for x in d.get("producao_artistica") or []:
        out.append(Item("artistica", ano_int(x.get("ano")), x.get("titulo"), x.get("autores")))
    ori = d.get("orientacoes") or {}
    for k, _r, chave in NIVEIS:
        for x in (ori.get("concluidas") or {}).get(chave) or []:
            out.append(Item("oc_" + k, ano_int(x.get("ano_conclusao") or x.get("ano")), x.get("titulo"),
                            x.get("orientando"), x.get("instituicao"), extra=x.get("tipo_orientacao", "")))
        for x in (ori.get("em_andamento") or {}).get(chave) or []:
            out.append(Item("oa_" + k, ano_int(x.get("ano_inicio") or x.get("ano")), x.get("titulo"),
                            x.get("orientando"), x.get("instituicao")))
    for chave, lista in d.items():
        if not chave.startswith("projetos_") or not isinstance(lista, list):
            continue
        cat = {"projetos_pesquisa": "pj_pesquisa", "projetos_extensao": "pj_extensao",
               "projetos_desenvolvimento": "pj_desenvolvimento"}.get(chave, "pj_outros")
        for x in lista:
            fim = str(x.get("ano_conclusao") or "")
            integ = x.get("integrantes") or []
            coord = [i.get("nome", "") for i in integ if "coord" in norm(i.get("papel"))]
            out.append(Item(cat, ano_int(x.get("ano_inicio")), x.get("nome"), "; ".join(i.get("nome", "") for i in integ),
                            fim=None if not ano_int(fim) else ano_int(fim), coord=coord,
                            extra="Em andamento" if not ano_int(fim) else "Concluído"))
    for x in d.get("premios_titulos") or []:
        out.append(Item("premios", ano_int(x.get("ano")), x.get("descricao")))
    for _nivel, lista in (d.get("bancas") or {}).items():
        for x in lista or []:
            out.append(Item("bancas", ano_int(x.get("ano")), x.get("descricao"), x.get("aluno"), extra=x.get("tipo", "")))
    ev = d.get("eventos") or {}
    for x in ev.get("participacoes") or []:
        out.append(Item("ev_part", ano_int(x.get("ano")), x.get("apresentacao") or x.get("evento"), veiculo=x.get("evento"),
                        extra=x.get("tipo_evento", "")))
    for x in ev.get("organizacoes") or []:
        out.append(Item("ev_org", ano_int(x.get("ano")), x.get("evento") or x.get("autores"), extra=x.get("natureza", "")))
    return out


class Base:
    """Itens do grupo inteiro, sem repeticao: o mesmo trabalho no CV de dois
    membros conta uma vez para o grupo (e uma vez para cada um deles)."""

    def __init__(self, grupo):
        self.grupo = grupo
        self.pessoas = {}   # id -> {"nome","atualizado"}
        self.itens = []
        unidos = {}
        for arq in sorted(glob.glob(os.path.join(grupo.saida, "json", "*.json"))):
            d = ler_json(arq, None)
            if not d:
                continue
            info = d.get("informacoes_pessoais") or {}
            lid = str(info.get("id_lattes") or "")
            if not lid:
                continue
            self.pessoas[lid] = {"nome": info.get("nome_completo") or lid, "atualizado": info.get("atualizacao_cv", "")}
            for it in itens_do_json(d):
                if not it.titulo:
                    continue
                k = it.chave(lid)
                if k in unidos:
                    velho = unidos[k]
                    velho.membros.add(lid)
                    velho.coord |= it.coord
                    if not velho.ano and it.ano:
                        velho.ano = it.ano
                else:
                    it.membros.add(lid)
                    unidos[k] = it
        self.itens = list(unidos.values())
        anos = [it.ano for it in self.itens if it.ano]
        self.ano_min = min(anos) if anos else datetime.date.today().year
        self.ano_max = max(max(anos) if anos else 0, datetime.date.today().year)


# ----------------------------------------------------------------- calculo

class Estat:
    def __init__(self, base, sel, a, b):
        self.a, self.b, self.sel = a, b, set(sel)
        self.anos = list(range(a, b + 1))
        nomes_sel = {norm(base.pessoas[i]["nome"]) for i in sel if i in base.pessoas}
        self.itens = []
        for it in base.itens:
            quem = it.membros & self.sel
            if not quem or not self._no_periodo(it):
                continue
            self.itens.append((it, quem))
        self.n = collections.Counter(it.cat for it, _q in self.itens)
        self.por_ano = collections.defaultdict(collections.Counter)
        for it, _q in self.itens:
            if it.cat in PJ:
                for y in self.anos:
                    if (it.ano or a) <= y and (it.fim is None or it.fim >= y):
                        self.por_ano["pj_ativos"][y] += 1
            elif not it.cat.startswith("oa_") and it.ano:
                self.por_ano[it.cat][it.ano] += 1
        self.coordenados = sum(1 for it, _q in self.itens if it.cat in PJ and {norm(c) for c in it.coord} & nomes_sel)
        artigos = [q for it, q in self.itens if it.cat == "artigos"]
        self.artigos_parceria = sum(1 for q in artigos if len(q) >= 2)
        self.por_membro = collections.defaultdict(collections.Counter)
        for it, quem in self.itens:
            for p in quem:
                self.por_membro[p][it.cat] += 1
        self.artigos_doi = sum(1 for it, _q in self.itens if it.cat == "artigos" and it.doi)
        self.anterior = None
        # Coautoria dentro do grupo: pares de membros que dividem um trabalho.
        self.pares = collections.Counter()
        for it, quem in self.itens:
            if it.cat in PRODUCAO and len(quem) >= 2:
                q = sorted(quem)
                for i in range(len(q)):
                    for j in range(i + 1, len(q)):
                        self.pares[(q[i], q[j])] += 1

    def _no_periodo(self, it):
        if it.cat.startswith("oa_"):
            return it.ano is None or it.ano <= self.b
        if it.cat in PJ:
            return (it.ano is None or it.ano <= self.b) and (it.fim is None or it.fim >= self.a)
        return it.ano is not None and self.a <= it.ano <= self.b

    def soma(self, cats):
        return sum(self.n[c] for c in cats)

    def de(self, cats):
        cats = set(cats)
        return [(it, q) for it, q in self.itens if it.cat in cats]


def rotulo(cat):
    return CATS[cat][0] if cat in CATS else cat


def veiculos(est, cats, limite=None):
    cont, nomes = collections.Counter(), collections.defaultdict(collections.Counter)
    for it, _q in est.de(cats):
        if it.veiculo:
            k = chave_texto(it.veiculo, 200)
            cont[k] += 1
            nomes[k][it.veiculo] += 1
    return [(nomes[k].most_common(1)[0][0], n) for k, n in cont.most_common(limite)]


# -------------------------------------------------------------- metricas

BIB_CATS = [k for k, _r, _j, _v in BIB]
DOCENTE = ("docente", "docentes", "professor", "professores", "professora", "professoras")


def com_anterior(base, sel, a, b):
    """Estatisticas do periodo e, em .anterior, as do periodo de mesma duracao logo antes."""
    est = Estat(base, sel, a, b)
    n = b - a + 1
    if a - n >= base.ano_min - 1:
        est.anterior = Estat(base, sel, a - n, a - 1)
    return est


def docentes(est, grupo):
    return [i for i in est.sel if norm((grupo.membro(i) or {}).get("categoria")) in DOCENTE]


def variacao(atual, antes):
    if not antes:
        return None
    return round(100 * (atual - antes) / antes)


def _dec(x, casas=1):
    return f"{x:.{casas}f}".replace(".", ",")


def indicadores(est, base, grupo):
    """[(nome, valor, detalhe)] das metricas do periodo."""
    anos = est.b - est.a + 1
    doc = docentes(est, grupo)
    nd = len(doc) or len(est.sel) or 1
    quem = "docente" if doc else "membro"
    bib = est.soma(BIB_CATS)
    ativos = sum(1 for i in est.sel if any(est.por_membro.get(i, {}).get(c) for c in BIB_CATS))
    n_art = est.n["artigos"]
    out = [
        ("Produção bibliográfica", num(bib), f"{_dec(bib / anos)} por ano"),
        ("Artigos em periódicos por ano", _dec(n_art / anos), f"{num(n_art)} no período"),
        (f"Artigos por {quem} por ano", _dec(n_art / nd / anos, 2),
         qt(len(doc), "membro na categoria Docentes", "membros na categoria Docentes") if doc else "não há membros na categoria Docentes"),
        (f"Orientações concluídas por {quem}", _dec(est.soma(OC) / nd), f"{num(est.soma(OC))} no período"),
        ("Colaboração interna", f"{round(100 * est.artigos_parceria / n_art)}%" if n_art else "–", "artigos com 2 ou mais membros do grupo"),
        ("Artigos com DOI", f"{round(100 * est.artigos_doi / n_art)}%" if n_art else "–", f"{num(est.artigos_doi)} artigos"),
        ("Membros com produção", f"{ativos} de {len(est.sel)}", "ao menos um item bibliográfico no período"),
        ("Periódicos distintos", num(len(veiculos(est, ["artigos"]))), "onde saíram os artigos"),
        ("Projetos coordenados pelo grupo", num(est.coordenados), f"de {num(est.soma(PJ))} projetos ativos"),
        ("Parcerias internas", num(len(est.pares)), "pares de membros com trabalho em comum"),
    ]
    ant = est.anterior
    if ant is not None:
        for nome, x, y in (("Artigos no período anterior", n_art, ant.n["artigos"]),
                           ("Produção bibliográfica no período anterior", bib, ant.soma(BIB_CATS)),
                           ("Orientações concluídas no período anterior", est.soma(OC), ant.soma(OC))):
            v = variacao(x, y)
            out.append((nome, num(y), f"{ant.a}–{ant.b}" + ("" if v is None else f" · variação de {'+' if v >= 0 else ''}{v}%")))
    return out


def tabela_indicadores(est, base, grupo):
    return ["Indicador", "Valor", "Detalhe"], [list(x) for x in indicadores(est, base, grupo)]


# -------------------------------------------------------------- graficos

# ------------------------------------------------------------- ficha

def ficha(est, base, grupo, lid):
    """Numeros de um membro no periodo, producao por ano e parcerias no grupo."""
    c = est.por_membro.get(lid, collections.Counter())
    por_ano = collections.defaultdict(collections.Counter)
    for it, quem in est.itens:
        if lid in quem and it.ano and it.cat in BIB_CATS:
            por_ano["artigos" if it.cat == "artigos" else "eventos" if it.cat in ("completos", "expandidos", "resumos") else "outros"][it.ano] += 1
    anos = est.anos
    spec = {"tipo": "colunas", "titulo": "Produção bibliográfica por ano", "categorias": anos, "series": [
        ("Artigos em periódicos", [por_ano["artigos"].get(y, 0) for y in anos]),
        ("Trabalhos em eventos", [por_ano["eventos"].get(y, 0) for y in anos]),
        ("Outros", [por_ano["outros"].get(y, 0) for y in anos])]}
    parceiros = collections.Counter()
    for (a, b), p in est.pares.items():
        if lid in (a, b):
            parceiros[b if a == lid else a] += p
    m = grupo.membro(lid) or {}
    return {
        "nome": base.pessoas.get(lid, {}).get("nome") or m.get("nome") or lid,
        "categoria": m.get("categoria", ""),
        "atualizado": base.pessoas.get(lid, {}).get("atualizado", ""),
        "numeros": [("Artigos", c["artigos"]), ("Eventos (completos)", c["completos"]), ("Livros e capítulos", c["livros"] + c["capitulos"]),
                    ("Orient. concluídas", sum(c[k] for k in OC)), ("Orient. em andamento", sum(c[k] for k in OA)), ("Projetos", sum(c[k] for k in PJ))],
        "grafico": spec,
        "parceiros": [(base.pessoas.get(i, {}).get("nome", i), p) for i, p in parceiros.most_common(12)],
    }


def iniciais(nome):
    partes = [x for x in re.split(r"\s+", str(nome or "").strip()) if x and x.lower() not in ("de", "da", "do", "dos", "das", "e")]
    return ((partes[0][0] + (partes[-1][0] if len(partes) > 1 else "")) if partes else "?").upper()


def graficos(est, base, grupo):
    """[(chave, titulo, spec)] na ordem do menu de graficos."""
    anos = est.anos
    pa = est.por_ano

    def serie(cats):
        return [sum(pa.get(c, {}).get(y, 0) for c in cats) for y in anos]

    out = [
        ("producao", "Produção bibliográfica por ano", {"tipo": "colunas", "categorias": anos, "series": [
            ("Artigos em periódicos", serie(["artigos"])),
            ("Trabalhos em eventos", serie(["completos", "expandidos", "resumos"])),
            ("Livros e capítulos", serie(["livros", "capitulos"])),
            ("Outros", serie(["aceitos", "jornais", "outras_bib", "apresentacoes"]))]}),
        ("artigos", "Artigos em periódicos por ano", {"tipo": "colunas", "categorias": anos, "series": [("Artigos", serie(["artigos"]))]}),
        ("orientacoes", "Orientações concluídas por ano", {"tipo": "colunas", "categorias": anos, "series": [
            ("Doutorado", serie(["oc_doutorado"])), ("Mestrado", serie(["oc_mestrado"])),
            ("Iniciação científica", serie(["oc_ic"])), ("TCC", serie(["oc_tcc"])),
            ("Outras", serie(["oc_posdoc", "oc_especializacao", "oc_outras"]))]}),
        ("projetos", "Projetos ativos por ano", {"tipo": "colunas", "categorias": anos, "series": [("Projetos", serie(["pj_ativos"]))]}),
        ("tecnica", "Produção técnica por ano", {"tipo": "colunas", "categorias": anos, "series": [
            ("Softwares", serie(["software"])), ("Patentes e registros", serie(["patentes"])), ("Outras", serie(["tecnica"]))]}),
    ]
    por = sorted(((base.pessoas[i]["nome"], est.por_membro.get(i, {}).get("artigos", 0)) for i in est.sel if i in base.pessoas),
                 key=lambda x: (-x[1], norm(x[0])))
    por = [x for x in por if x[1]][:15]
    out.append(("membros", "Artigos por membro (15 primeiros)", {"tipo": "barras", "rotulos": [x[0] for x in por], "valores": [x[1] for x in por]}))

    cats = collections.Counter()
    for _it, quem in est.de(BIB_CATS):
        for c in {(grupo.membro(p) or {}).get("categoria") or "Sem categoria" for p in quem}:
            cats[c] += 1
    cc = cats.most_common()
    out.append(("categorias", "Produção por categoria de membro", {"tipo": "barras", "rotulos": [c for c, _n in cc], "valores": [n for _c, n in cc]}))

    tipos = sorted([(rotulo(c), est.n[c]) for c in CATS if c not in OA and c not in PJ and est.n[c]], key=lambda x: -x[1])
    out.append(("tipos", "Itens por tipo no período", {"tipo": "barras", "rotulos": [t for t, _n in tipos], "valores": [n for _t, n in tipos]}))

    v = veiculos(est, ["artigos", "aceitos"], 12)
    out.append(("periodicos", "Periódicos mais frequentes", {"tipo": "barras", "rotulos": [x for x, _n in v], "valores": [n for _x, n in v]}))
    v = veiculos(est, ["completos", "expandidos", "resumos"], 12)
    out.append(("eventos", "Eventos mais frequentes", {"tipo": "barras", "rotulos": [x for x, _n in v], "valores": [n for _x, n in v]}))

    cc = collections.Counter((grupo.membro(i) or {}).get("categoria") or "Sem categoria" for i in est.sel).most_common()
    out.append(("composicao", "Composição do grupo", {"tipo": "barras", "rotulos": [c for c, _n in cc], "valores": [n for _c, n in cc], "cor": marca.SERIES[2]}))

    out.append(("rede", "Rede de colaboração entre os membros", rede(est, base, grupo)))
    for _k, titulo, spec in out:
        spec["titulo"] = titulo
    return out


def posicoes(pares):
    """Posicoes (0..1) dos membros na rede. Cada parte desconectada e arrumada
    a parte: a maior ocupa a area principal, as menores ficam numa coluna a
    direita (o spring_layout de tudo junto joga as partes pequenas para longe)."""
    try:
        import networkx as nx
    except ImportError:
        nos = sorted({i for p in pares for i in p})
        return {i: (0.5 + 0.5 * math.cos(2 * math.pi * k / len(nos)), 0.5 + 0.5 * math.sin(2 * math.pi * k / len(nos))) for k, i in enumerate(nos)}
    g = nx.Graph()
    for (a, b), p in pares.items():
        g.add_edge(a, b, weight=p)
    if not len(g):
        return {}
    partes = sorted(nx.connected_components(g), key=len, reverse=True)

    def arrumar(nos, x0, y0, x1, y1):
        sub = g.subgraph(nos)
        if len(nos) == 1:
            return {next(iter(nos)): ((x0 + x1) / 2, (y0 + y1) / 2)}
        if len(nos) == 2:
            a, b = sorted(nos)
            return {a: (x0 + (x1 - x0) * 0.25, (y0 + y1) / 2), b: (x0 + (x1 - x0) * 0.75, (y0 + y1) / 2)}
        try:
            bruto = nx.kamada_kawai_layout(sub, weight=None)
        except Exception:  # noqa: BLE001 - sem scipy
            bruto = nx.spring_layout(sub, seed=7, iterations=200)
        xs = [float(p[0]) for p in bruto.values()]
        ys = [float(p[1]) for p in bruto.values()]
        dx, dy = (max(xs) - min(xs)) or 1, (max(ys) - min(ys)) or 1
        return {i: (x0 + (float(p[0]) - min(xs)) / dx * (x1 - x0), y0 + (float(p[1]) - min(ys)) / dy * (y1 - y0)) for i, p in bruto.items()}

    pos = {}
    if len(partes) == 1:
        return arrumar(partes[0], 0, 0, 1, 1)
    pos.update(arrumar(partes[0], 0, 0, 0.74, 1))
    resto = partes[1:]
    alt = 1 / len(resto)
    for k, parte in enumerate(resto):
        pos.update(arrumar(parte, 0.82, k * alt + alt * 0.15, 1, (k + 1) * alt - alt * 0.15))
    return pos


def rede(est, base, grupo):
    """Coautoria no grupo: no = membro (tamanho = producao), aresta = trabalhos em comum."""
    nos_ids = {i for par in est.pares for i in par}
    prod = {i: sum(est.por_membro.get(i, {}).get(c, 0) for c in PRODUCAO) for i in est.sel}
    cats = collections.Counter((grupo.membro(i) or {}).get("categoria") or "Sem categoria" for i in nos_ids)
    grupos = [c for c, _n in cats.most_common()]
    if len(grupos) > 6:  # mais de 6 cores nao se distinguem: o resto vira "Outras"
        grupos = grupos[:5] + ["Outras"]
    pos = posicoes(est.pares)

    def grupo_de(i):
        c = (grupo.membro(i) or {}).get("categoria") or "Sem categoria"
        return c if c in grupos else "Outras"
    nos = [(i, base.pessoas.get(i, {}).get("nome", i), prod.get(i, 0), grupo_de(i)) for i in nos_ids]
    return {"tipo": "rede", "nos": nos, "arestas": [(a, b, p) for (a, b), p in est.pares.items()], "pos": pos, "grupos": grupos}


# ------------------------------------------------- tabelas (tela e Excel)

def tabela_por_ano(est):
    cab = ["Indicador"] + [str(y) for y in est.anos] + ["Total"]
    linhas = []
    for cat in [c for c in CATS if c not in OA and c not in PJ]:
        cont = est.por_ano.get(cat)
        if cont and sum(cont.values()):
            linhas.append([rotulo(cat) + (" (concluídas)" if cat in OC else "")] + [cont.get(y, 0) for y in est.anos] + [sum(cont.values())])
    pj = est.por_ano.get("pj_ativos")
    if pj:
        linhas.append(["Projetos ativos no ano"] + [pj.get(y, 0) for y in est.anos] + [est.soma(PJ)])
    return cab, linhas


def tabela_por_membro(est, base, grupo):
    cab = ["Membro", "Categoria", "Artigos", "Eventos (completos)", "Livros e capítulos", "Produção técnica",
           "Orient. concluídas", "Orient. em andamento", "Projetos", "Prêmios", "Bancas", "CV atualizado em"]
    linhas = []
    for lid in est.sel:
        if lid not in base.pessoas:
            continue
        c = est.por_membro.get(lid, collections.Counter())
        m = grupo.membro(lid) or {}
        linhas.append([base.pessoas[lid]["nome"], m.get("categoria", ""), c["artigos"], c["completos"],
                       c["livros"] + c["capitulos"], c["software"] + c["patentes"] + c["tecnica"],
                       sum(c[k] for k in OC), sum(c[k] for k in OA), sum(c[k] for k in PJ), c["premios"], c["bancas"],
                       base.pessoas[lid]["atualizado"]])
    linhas.sort(key=lambda r: (-r[2], -r[3], norm(r[0])))
    return cab, linhas


def tabela_orientacoes_resumo(est):
    cab = ["Nível", "Concluídas no período", "Em andamento"]
    linhas = [[r, est.n["oc_" + k], est.n["oa_" + k]] for k, r, _j in NIVEIS]
    linhas.append(["Total", est.soma(OC), est.soma(OA)])
    return cab, linhas


def tabela_orientacoes(est, base):
    cab = ["Situação", "Nível", "Ano", "Orientando", "Título", "Orientador(es) do grupo"]
    linhas = []
    for it, q in est.de(OC + OA):
        linhas.append(["Concluída" if it.cat in OC else "Em andamento", rotulo(it.cat), it.ano or "", it.pessoas, it.titulo,
                       "; ".join(sorted(base.pessoas[p]["nome"] for p in q))])
    linhas.sort(key=lambda r: (r[0] != "Em andamento", -(r[2] or 0), r[1]))
    return cab, linhas


def tabela_projetos(est, base):
    cab = ["Natureza", "Início", "Fim", "Projeto", "Coordenação", "Membros do grupo"]
    linhas = []
    for it, q in est.de(PJ):
        linhas.append([rotulo(it.cat).replace("Projetos de ", "").replace("Outros projetos", "outra").capitalize(), it.ano or "",
                       it.fim or "atual", it.titulo, "; ".join(sorted(it.coord)), "; ".join(sorted(base.pessoas[p]["nome"] for p in q))])
    linhas.sort(key=lambda r: (r[2] != "atual", -(r[1] or 0), norm(r[3])))
    return cab, linhas


def tabela_veiculos(est):
    cab = ["Tipo", "Periódico ou evento", "Trabalhos"]
    linhas = [["Periódico", v, n] for v, n in veiculos(est, ["artigos", "aceitos"])]
    linhas += [["Evento", v, n] for v, n in veiculos(est, ["completos", "expandidos", "resumos"])]
    return cab, linhas


def tabela_itens(est, base, cats=None, busca=""):
    cab = ["Tipo", "Ano", "Título", "Autores / pessoas", "Veículo", "DOI", "Membros do grupo"]
    termos = norm(busca).split()
    linhas = []
    for it, q in est.itens:
        if cats and it.cat not in cats:
            continue
        nomes = "; ".join(sorted(base.pessoas[p]["nome"] for p in q))
        if termos:
            alvo = norm(" ".join((it.titulo, it.pessoas, it.veiculo, it.doi, nomes)))
            if not all(t in alvo for t in termos):
                continue
        linhas.append([rotulo(it.cat), it.ano or "", it.titulo, it.pessoas, it.veiculo, it.doi, nomes])
    linhas.sort(key=lambda r: (-(r[1] or 0), list(CATS).index(next(k for k in CATS if CATS[k][0] == r[0])), norm(r[2])))
    return cab, linhas


def texto_relatorio(est, grupo, base):
    a, b = est.a, est.b
    per = f"em {a}" if a == b else f"de {a} a {b}"
    n_memb = len([i for i in est.sel if i in base.pessoas])
    cats = collections.Counter((grupo.membro(i) or {}).get("categoria") or "Sem categoria" for i in est.sel if i in base.pessoas)
    comp = lista_pt([f"{n} {c.lower()}" for c, n in cats.most_common()])
    p = []
    p.append(f"Este levantamento considera {qt(n_memb, 'membro', 'membros')} do grupo {grupo.nome}" + (f" ({comp})" if comp else "")
             + f" e o período {per}, com base nos Currículos Lattes. Trabalhos que aparecem no currículo de mais de um membro foram contados uma única vez.")

    bib = []
    for cat, sing, plur in (("artigos", "artigo em periódico", "artigos em periódicos"),
                            ("aceitos", "artigo aceito para publicação", "artigos aceitos para publicação"),
                            ("livros", "livro", "livros"), ("capitulos", "capítulo de livro", "capítulos de livros"),
                            ("completos", "trabalho completo em anais de eventos", "trabalhos completos em anais de eventos"),
                            ("expandidos", "resumo expandido", "resumos expandidos"), ("resumos", "resumo em eventos", "resumos em eventos")):
        if est.n[cat]:
            bib.append(qt(est.n[cat], sing, plur))
    if bib:
        t = f"No período, o grupo publicou {lista_pt(bib)}."
        if est.n["artigos"] >= 2:
            pct = round(100 * est.artigos_parceria / est.n["artigos"])
            t += f" Dos artigos em periódicos, {num(est.artigos_parceria)} ({pct}%) têm dois ou mais membros do grupo entre os autores."
        top = veiculos(est, ["artigos"], 3)
        if top and est.n["artigos"] >= 5:
            t += " Os periódicos mais frequentes foram " + lista_pt([f"{v} ({n})" for v, n in top]) + "."
        p.append(t)
    else:
        p.append("Não há produção bibliográfica registrada no período.")

    ant = est.anterior
    if ant is not None and ant.n["artigos"]:
        v = variacao(est.n["artigos"], ant.n["artigos"])
        rumo = "cresceu" if v > 0 else "caiu" if v < 0 else "se manteve"
        p.append(f"Em relação ao período anterior de mesma duração ({ant.a}–{ant.b}), quando foram publicados {qt(ant.n['artigos'], 'artigo', 'artigos')} "
                 f"em periódicos, o número de artigos {rumo}" + (f" {abs(v)}%." if v else "."))
    doc = docentes(est, grupo)
    if doc and est.n["artigos"]:
        media = f"{est.n['artigos'] / len(doc) / (b - a + 1):.2f}".replace(".", ",")
        p.append(f"A média foi de {media} artigos por docente por ano ({qt(len(doc), 'docente', 'docentes')} no grupo).")

    tec = []
    for cat, sing, plur in (("software", "software", "softwares"), ("patentes", "patente ou registro", "patentes e registros"),
                            ("tecnica", "outra produção técnica", "outras produções técnicas")):
        if est.n[cat]:
            tec.append(qt(est.n[cat], sing, plur))
    if tec:
        p.append(f"Na produção técnica, há registro de {lista_pt(tec)}.")

    nomes_niv = {"posdoc": ("supervisão de pós-doutorado", "supervisões de pós-doutorado"), "doutorado": ("tese de doutorado", "teses de doutorado"),
                 "mestrado": ("dissertação de mestrado", "dissertações de mestrado"), "especializacao": ("monografia de especialização", "monografias de especialização"),
                 "tcc": ("trabalho de conclusão de curso", "trabalhos de conclusão de curso"), "ic": ("iniciação científica", "iniciações científicas"),
                 "outras": ("orientação de outra natureza", "orientações de outra natureza")}
    conc = [qt(est.n["oc_" + k], *nomes_niv[k]) for k, _r, _j in NIVEIS if est.n["oc_" + k]]
    anda = [qt(est.n["oa_" + k], *nomes_niv[k]) for k, _r, _j in NIVEIS if est.n["oa_" + k]]
    if conc or anda:
        t = ""
        if conc:
            t += f"Na formação de recursos humanos, foram concluídas {num(est.soma(OC))} orientações: {lista_pt(conc)}."
        if anda:
            t += (" " if t else "") + f"Estão em andamento {num(est.soma(OA))} orientações ({lista_pt(anda)})."
        p.append(t)

    if est.soma(PJ):
        tipos = lista_pt([qt(est.n[c], *{"pj_pesquisa": ("de pesquisa", "de pesquisa"), "pj_extensao": ("de extensão", "de extensão"),
                                            "pj_desenvolvimento": ("de desenvolvimento", "de desenvolvimento"),
                                            "pj_outros": ("de outra natureza", "de outra natureza")}[c]) for c in PJ if est.n[c]])
        t = f"Os membros participaram de {qt(est.soma(PJ), 'projeto', 'projetos')} ativo{'s' if est.soma(PJ) != 1 else ''} no período ({tipos})"
        t += f", {num(est.coordenados)} deles coordenado{'s' if est.coordenados != 1 else ''} por membros do grupo." if est.coordenados else "."
        p.append(t)

    outros = []
    if est.n["premios"]:
        outros.append(qt(est.n["premios"], "prêmio ou título", "prêmios e títulos"))
    if est.n["bancas"]:
        outros.append(qt(est.n["bancas"], "participação em banca", "participações em bancas"))
    if est.n["ev_part"]:
        outros.append(qt(est.n["ev_part"], "participação em evento", "participações em eventos"))
    if est.n["ev_org"]:
        outros.append(qt(est.n["ev_org"], "evento organizado", "eventos organizados"))
    if outros:
        p.append(f"Há registro ainda de {lista_pt(outros)}.")
    return "\n\n".join(p)


def salvar_excel(caminho, abas):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    wb.remove(wb.active)
    for nome, cab, linhas in abas:
        ws = wb.create_sheet(nome[:31])
        ws.append(cab)
        for r in linhas:
            ws.append(r)
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1F5FA8")
            c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.freeze_panes = "B2" if len(cab) > 3 else "A2"
        for i, titulo in enumerate(cab, 1):
            maior = max([len(str(titulo))] + [len(str(r[i - 1])) for r in linhas[:500] if i - 1 < len(r)])
            ws.column_dimensions[get_column_letter(i)].width = min(max(9, maior + 2), 70)
        if linhas:
            ws.auto_filter.ref = ws.dimensions
    wb.save(caminho)


